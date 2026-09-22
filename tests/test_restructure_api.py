"""API tests for the Excel restructure endpoints (no LLM: plans are passed in)."""

from __future__ import annotations

import io
from typing import Any

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook

from extractor.api.main import app
from extractor.tabular import ColumnOp, ColumnOpKind, TablePlan


def xlsx_bytes(rows: list[list[Any]], sheet: str = "Sheet1") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def simple_plan() -> TablePlan:
    return TablePlan(
        columns=[
            ColumnOp(op=ColumnOpKind.rename, column="A", target="item"),
            ColumnOp(op=ColumnOpKind.coerce_number, column="C"),
        ]
    )


@pytest.fixture
def client() -> AsyncClient:
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


async def test_apply_returns_rows_and_stats(client: AsyncClient) -> None:
    data = xlsx_bytes([["A", "B", "C"], ["x", "y", "10"], ["p", "q", "2"]])
    async with client as c:
        resp = await c.post(
            "/api/restructure/apply",
            files={"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"plan": TablePlan(columns=[ColumnOp(op=ColumnOpKind.coerce_number, column="C")]).model_dump_json()},
        )
    assert resp.status_code == 200
    body = resp.json()
    assert body["columns"] == ["A", "B", "C"]
    assert body["rows"][0][2] == 10.0
    assert body["stats"]["rows_out"] == 2


async def test_apply_rename_updates_columns(client: AsyncClient) -> None:
    data = xlsx_bytes([["A", "B", "C"], ["x", "y", "5"]])
    async with client as c:
        resp = await c.post(
            "/api/restructure/apply",
            files={"file": ("t.xlsx", data, "application/octet-stream")},
            data={"plan": simple_plan().model_dump_json()},
        )
    assert resp.status_code == 200
    assert resp.json()["columns"] == ["item", "B", "C"]


async def test_download_returns_xlsx_bytes(client: AsyncClient) -> None:
    data = xlsx_bytes([["A"], ["hello"]])
    async with client as c:
        resp = await c.post(
            "/api/restructure/download",
            files={"file": ("t.xlsx", data, "application/octet-stream")},
            data={"plan": TablePlan().model_dump_json()},
        )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    # round-trip: re-open what the server produced
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(resp.content))
    assert list(next(wb.active.iter_rows(values_only=True))) == ["A"]


async def test_sheet_selection_and_404(client: AsyncClient) -> None:
    data = xlsx_bytes([["A"], [1]], sheet="Data")
    async with client as c:
        ok = await c.post(
            "/api/restructure/apply",
            files={"file": ("t.xlsx", data, "application/octet-stream")},
            data={"plan": TablePlan().model_dump_json(), "sheet": "data"},
        )
        missing = await c.post(
            "/api/restructure/apply",
            files={"file": ("t.xlsx", data, "application/octet-stream")},
            data={"plan": TablePlan().model_dump_json(), "sheet": "nope"},
        )
    assert ok.status_code == 200
    assert ok.json()["sheet_name"] == "Data"
    assert missing.status_code == 404


async def test_unknown_plan_column_maps_to_422(client: AsyncClient) -> None:
    data = xlsx_bytes([["A"], [1]])
    async with client as c:
        resp = await c.post(
            "/api/restructure/apply",
            files={"file": ("t.xlsx", data, "application/octet-stream")},
            data={"plan": TablePlan(columns=[ColumnOp(op=ColumnOpKind.trim, column="Zed")]).model_dump_json()},
        )
    assert resp.status_code == 422


async def test_invalid_plan_json_maps_to_400(client: AsyncClient) -> None:
    data = xlsx_bytes([["A"], [1]])
    async with client as c:
        resp = await c.post(
            "/api/restructure/apply",
            files={"file": ("t.xlsx", data, "application/octet-stream")},
            data={"plan": "{not json"},
        )
    assert resp.status_code == 400


async def test_plan_template_and_health(client: AsyncClient) -> None:
    async with client as c:
        template = await c.get("/api/plan-template")
        health = await c.get("/api/health")
    assert template.status_code == 200
    assert "concat" in template.json()["ops"]
    assert health.status_code == 200
    assert "product" in health.json()["schemas"]
